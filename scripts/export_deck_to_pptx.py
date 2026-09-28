#!/usr/bin/env python3
"""
Export HTML presentation (1920x1080 fixed stage) to 16:9 PPTX format
with Retina-quality resolution and speaker notes.
"""
import os
import sys
import time
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

def export_html_to_pptx(html_path, output_pptx_path):
    abs_html = os.path.abspath(html_path)
    if not os.path.exists(abs_html):
        print(f"Error: {abs_html} not found.")
        sys.exit(1)

    print(f"Converting: {abs_html}")
    print(f"Target PPTX: {output_pptx_path}")

    temp_dir = "/tmp/pptx_slides_export"
    os.makedirs(temp_dir, exist_ok=True)

    slide_images = []
    slide_notes = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        # Device scale factor 2 for crisp 3840x2160 rendering
        context = browser.new_context(
            viewport={"width": 1920, "height": 1080},
            device_scale_factor=2
        )
        page = context.new_page()

        file_url = f"file://{abs_html}"
        page.goto(file_url, wait_until="networkidle")

        # Wait for Google Fonts to finish loading
        page.wait_for_timeout(1500)

        # Hide UI controls
        page.evaluate("""
            const style = document.createElement('style');
            style.innerHTML = `
                .nav-controls, .slide-counter, .deck-top-actions, .edit-banner, .deck-controls, .edit-toggle, .edit-hotzone {
                    display: none !important;
                }
            `;
            document.head.appendChild(style);
        """)

        # Get total slides
        total_slides = page.evaluate("() => document.querySelectorAll('.slide').length")
        print(f"Detected {total_slides} slides.")

        for i in range(total_slides):
            # Navigate to slide i
            page.evaluate(f"() => {{ if (window.deck && window.deck.goTo) window.deck.goTo({i}); }}")
            page.wait_for_timeout(400)

            # Extract text content for speaker notes
            note_text = page.evaluate(f"""() => {{
                const s = document.querySelectorAll('.slide')[{i}];
                if (!s) return '';
                return s.innerText.trim();
            }}""")

            img_path = os.path.join(temp_dir, f"slide_{i+1:02d}.png")
            page.screenshot(path=img_path)
            slide_images.append(img_path)
            slide_notes.append(note_text)
            print(f"Captured slide {i+1}/{total_slides}")

        browser.close()

    # Build PPTX using python-pptx
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    for idx, (img_path, note) in enumerate(zip(slide_images, slide_notes)):
        slide = prs.slides.add_slide(blank_layout)
        # Add high-resolution image filling entire slide
        slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

        # Set speaker notes
        if note:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = note

        print(f"Added slide {idx+1} to PPTX")

    os.makedirs(os.path.dirname(os.path.abspath(output_pptx_path)), exist_ok=True)
    prs.save(output_pptx_path)
    print(f"Successfully generated: {output_pptx_path}")

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "/Users/ringi/ai-infra-odyssey/slides/ai-infra-odyssey-deck.html"
    dst = sys.argv[2] if len(sys.argv) > 2 else "/Users/ringi/ai-infra-odyssey/slides/AI_Infra_Odyssey_BlockFrame.pptx"
    export_html_to_pptx(src, dst)
